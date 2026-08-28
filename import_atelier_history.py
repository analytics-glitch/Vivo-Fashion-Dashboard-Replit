#!/usr/bin/env python3
"""Dry-run-first importer for the historical Junction Atelier tracker.

Parsing and validation deliberately have no database dependency.  Running the
script without ``--commit`` only prints a JSON reconciliation report.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import sys
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Mapping

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


DEFAULT_WORKBOOK = (
    "attached_assets/"
    "0_2026_-_Vivo_Junction_In-store_Tailoring___Adjustments_Log_"
    "(_1787907179635.xlsx"
)
SOURCE_SYSTEM = "junction-atelier-tracker"
SOURCE_SHEET = "Tracker"
EARLIEST_LOG_DATE = date(2026, 6, 1)
LATEST_LOG_DATE = date(2026, 8, 31)
NULL_MARKERS = {"", "-", "missing"}

HEADERS = {
    "date_in": "Date In",
    "customer_name": "Customer Name",
    "phone": "Phone no",
    "service_type": "Service",
    "sku": "Item Code",
    "item_description": "Description (as per system)",
    "tailor_name": "Tailor's Name",
    "qty": "Qty",
    "alteration_detail": "Alteration",
    "promised_date": "Collection Date",
    "status": "Status",
    "date_out": "Date Out",
    "notes": "Notes",
}

STATUS_MAP = {
    "pending": "Received",
    "received": "Received",
    "in progress": "In Progress",
    "in-progress": "In Progress",
    "ready": "Ready for Pickup",
    "ready for pickup": "Ready for Pickup",
    "collected": "Collected",
    "complete": "Collected",
    "completed": "Collected",
    "cancelled": "Cancelled",
    "canceled": "Cancelled",
}


@dataclass(frozen=True)
class ImportJob:
    source_identity: str
    source_row: int
    source_item: int
    customer_name: str | None
    phone: str | None
    service_type: str | None
    sku: str | None
    item_description: str | None
    staff_name: str | None
    alteration_detail: str | None
    status: str
    date_in: datetime
    promised_date: datetime | None
    date_ready: datetime | None
    collected_at: datetime | None
    notes: str | None
    qty: int = 1
    charge_amount: Decimal = Decimal("0")


@dataclass(frozen=True)
class Quarantine:
    source_row: int
    reasons: tuple[str, ...]
    expanded_quantity: int


@dataclass
class ParseResult:
    jobs: list[ImportJob]
    quarantined: list[Quarantine]
    source_rows: int
    blank_rows_skipped: int
    quantities: int
    invalid_phones: int
    source_sha256: str

    def report(self, *, committed: bool = False, inserted: int = 0, existing: int = 0) -> dict[str, Any]:
        statuses = Counter(job.status for job in self.jobs)
        staff = Counter(job.staff_name or "(unassigned)" for job in self.jobs)
        quarantined_jobs = sum(row.expanded_quantity for row in self.quarantined)
        return {
            "mode": "commit" if committed else "dry-run",
            "source": {
                "system": SOURCE_SYSTEM,
                "sheet": SOURCE_SHEET,
                "sha256": self.source_sha256,
            },
            "reconciliation": {
                "source_rows": self.source_rows,
                "source_quantity": self.quantities,
                "importable_source_rows": self.source_rows - len(self.quarantined),
                "importable_jobs": len(self.jobs),
                "quarantined_source_rows": len(self.quarantined),
                "quarantined_jobs": quarantined_jobs,
                "blank_rows_skipped": self.blank_rows_skipped,
                "invalid_phones_nullified": self.invalid_phones,
                "inserted": inserted,
                "already_present": existing,
            },
            "review": {
                "quarantine": [asdict(row) for row in self.quarantined],
                "status_totals": dict(sorted(statuses.items())),
                "staff_totals": dict(sorted(staff.items())),
            },
        }


def clean_text(value: Any) -> str | None:
    """Trim a cell and turn the workbook's null markers into ``None``."""
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    text = str(value).strip()
    return None if text.casefold() in NULL_MARKERS else text


def preserve_sku(value: Any) -> str | None:
    """Return SKU text, avoiding the ``.0`` added by integral Excel numbers."""
    if value is None or isinstance(value, bool):
        return clean_text(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, (float, Decimal)):
        try:
            number = Decimal(str(value))
        except InvalidOperation:
            return clean_text(value)
        if not number.is_finite():
            return None
        return str(int(number)) if number == number.to_integral() else format(number, "f")
    return clean_text(value)


def normalize_kenyan_phone(value: Any) -> str | None:
    """Normalize a Kenyan mobile number to ``+254XXXXXXXXX``.

    Excel frequently stores local numbers numerically and consequently drops
    the leading zero; both that representation and common international forms
    are accepted. Invalid/non-Kenyan values are nullified, not guessed.
    """
    text = preserve_sku(value)
    if text is None:
        return None
    digits = re.sub(r"\D", "", text)
    if digits.startswith("00254"):
        digits = digits[2:]
    if digits.startswith("254"):
        national = digits[3:]
    elif digits.startswith("0"):
        national = digits[1:]
    else:
        national = digits
    if len(national) == 9 and national[0] in {"1", "7"}:
        return f"+254{national}"
    return None


def map_tailor_name(value: Any) -> str | None:
    name = clean_text(value)
    if name is None:
        return None
    compact = re.sub(r"[^a-z]", "", name.casefold())
    if compact.startswith("maxi") or compact == "maxmilla":
        return "Maximillia Kubochi"
    return " ".join(part.capitalize() for part in name.split())


def map_legacy_status(value: Any) -> str:
    status = clean_text(value)
    if status is None:
        return "Received"
    key = re.sub(r"\s+", " ", status.casefold())
    if key not in STATUS_MAP:
        raise ValueError(f"unknown status {status!r}")
    return STATUS_MAP[key]


def parse_quantity(value: Any) -> int:
    if clean_text(value) is None:
        return 1
    try:
        quantity = Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"invalid quantity {value!r}") from exc
    if not quantity.is_finite() or quantity != quantity.to_integral() or quantity < 1:
        raise ValueError(f"invalid quantity {value!r}")
    return int(quantity)


def parse_date_cell(value: Any, field: str) -> datetime | None:
    if clean_text(value) is None:
        return None
    parsed: datetime | date
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = value
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            parsed = from_excel(value)
        except (TypeError, ValueError, OverflowError) as exc:
            raise ValueError(f"{field} is unparseable: {value!r}") from exc
    else:
        text = str(value).strip()
        parsed = None  # type: ignore[assignment]
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d %b %Y", "%d %B %Y"):
            try:
                parsed = datetime.strptime(text, fmt)
                break
            except ValueError:
                pass
        if parsed is None:
            raise ValueError(f"{field} is unparseable: {value!r}")
    result = datetime.combine(parsed, time.min) if type(parsed) is date else parsed
    if not EARLIEST_LOG_DATE <= result.date() <= LATEST_LOG_DATE:
        raise ValueError(f"{field} is outside log range: {result.date().isoformat()}")
    return result.replace(tzinfo=None)


def _source_identity(row_number: int, item_number: int) -> str:
    raw = f"{SOURCE_SYSTEM}:{SOURCE_SHEET}:{row_number}:{item_number}"
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _meaningful(values: Mapping[str, Any]) -> bool:
    return any(clean_text(value) is not None for value in values.values())


def parse_tracker_row(
    values: Mapping[str, Any], row_number: int
) -> tuple[list[ImportJob], Quarantine | None, bool]:
    """Parse one mapped Tracker row.

    The final boolean says that a non-empty, invalid phone was nullified.
    """
    try:
        quantity = parse_quantity(values.get("qty"))
    except ValueError as exc:
        return [], Quarantine(row_number, (str(exc),), 1), False

    errors: list[str] = []
    dates: dict[str, datetime | None] = {}
    for field in ("date_in", "promised_date", "date_out"):
        try:
            dates[field] = parse_date_cell(values.get(field), field)
        except ValueError as exc:
            errors.append(str(exc))
            dates[field] = None
    if dates["date_in"] is None and not any(reason.startswith("date_in ") for reason in errors):
        errors.append("date_in is required")
    if dates["date_in"]:
        for field in ("promised_date", "date_out"):
            if dates[field] and dates[field] < dates["date_in"]:
                errors.append(f"{field} is earlier than date_in")

    try:
        status = map_legacy_status(values.get("status"))
    except ValueError as exc:
        errors.append(str(exc))
        status = "Received"
    if errors:
        return [], Quarantine(row_number, tuple(errors), quantity), False

    raw_phone = clean_text(values.get("phone"))
    phone = normalize_kenyan_phone(values.get("phone"))
    invalid_phone = raw_phone is not None and phone is None
    date_out = dates["date_out"]
    date_ready = date_out if status in {"Ready for Pickup", "Collected"} else None
    collected_at = date_out if status == "Collected" else None

    common = {
        "source_row": row_number,
        "customer_name": clean_text(values.get("customer_name")),
        "phone": phone,
        "service_type": clean_text(values.get("service_type")),
        "sku": preserve_sku(values.get("sku")),
        "item_description": clean_text(values.get("item_description")),
        "staff_name": map_tailor_name(values.get("tailor_name")),
        "alteration_detail": clean_text(values.get("alteration_detail")),
        "status": status,
        "date_in": dates["date_in"],
        "promised_date": dates["promised_date"],
        "date_ready": date_ready,
        "collected_at": collected_at,
        "notes": clean_text(values.get("notes")),
    }
    jobs = [
        ImportJob(
            source_identity=_source_identity(row_number, item_number),
            source_item=item_number,
            **common,  # type: ignore[arg-type]
        )
        for item_number in range(1, quantity + 1)
    ]
    return jobs, None, invalid_phone


def parse_workbook(path: str | os.PathLike[str]) -> ParseResult:
    workbook_path = Path(path)
    source_hash = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    if SOURCE_SHEET not in workbook.sheetnames:
        raise ValueError(f"workbook has no {SOURCE_SHEET!r} sheet")
    sheet = workbook[SOURCE_SHEET]
    rows = sheet.iter_rows(values_only=True)
    try:
        raw_headers = next(rows)
    except StopIteration as exc:
        raise ValueError("Tracker sheet is empty") from exc
    header_positions = {clean_text(value): index for index, value in enumerate(raw_headers)}
    missing = [label for label in HEADERS.values() if label not in header_positions]
    if missing:
        raise ValueError(f"Tracker is missing required columns: {', '.join(missing)}")

    jobs: list[ImportJob] = []
    quarantined: list[Quarantine] = []
    source_rows = blank_rows = total_quantity = invalid_phones = 0
    for row_number, row in enumerate(rows, start=2):
        values = {
            key: row[header_positions[label]] if header_positions[label] < len(row) else None
            for key, label in HEADERS.items()
        }
        if not _meaningful(values):
            blank_rows += 1
            continue
        source_rows += 1
        try:
            quantity = parse_quantity(values["qty"])
        except ValueError:
            quantity = 1
        total_quantity += quantity
        parsed, quarantine, invalid_phone = parse_tracker_row(values, row_number)
        jobs.extend(parsed)
        if quarantine:
            quarantined.append(quarantine)
        invalid_phones += int(invalid_phone)
    workbook.close()
    return ParseResult(
        jobs=jobs,
        quarantined=quarantined,
        source_rows=source_rows,
        blank_rows_skipped=blank_rows,
        quantities=total_quantity,
        invalid_phones=invalid_phones,
        source_sha256=source_hash,
    )


def commit_jobs(result: ParseResult, database_url: str) -> tuple[int, int]:
    """Transactionally write jobs using the schema established by ``atelier.py``.

    A small import ledger owns source identity instead of overloading a
    customer-facing field. Its primary key plus the transaction advisory lock
    make concurrent and later reruns safe. Additive snapshot columns retain
    sparse legacy customer data without creating a second customer table.
    """
    try:
        import psycopg2
    except ImportError as exc:
        raise RuntimeError("psycopg2 is required with --commit") from exc

    db_status = {
        "Received": "intake",
        "In Progress": "in_progress",
        "Ready for Pickup": "ready",
        "Collected": "collected",
        "Cancelled": "cancelled",
    }
    inserted = existing = 0
    connection = psycopg2.connect(database_url)
    try:
        with connection:
            with connection.cursor() as cursor:
                cursor.execute(
                    "SELECT to_regclass('atelier_jobs'), "
                    "to_regclass('atelier_locations'), to_regclass('app_users')"
                )
                required = cursor.fetchone()
                if not required or any(item is None for item in required):
                    raise RuntimeError(
                        "Atelier tables are unavailable; deploy atelier.py before committing"
                    )
                cursor.execute(
                    "SELECT pg_advisory_xact_lock(hashtext(%s))", (SOURCE_SYSTEM,)
                )
                cursor.execute(
                    """SELECT id FROM atelier_locations
                       WHERE code='junction' AND active=TRUE ORDER BY id LIMIT 1"""
                )
                location = cursor.fetchone()
                if not location:
                    raise RuntimeError("active Junction Atelier location was not found")
                location_id = location[0]
                cursor.execute(
                    """SELECT user_id FROM app_users
                       WHERE status='active'
                       ORDER BY CASE WHEN lower(role)='admin' THEN 0 ELSE 1 END,
                                created_at, user_id LIMIT 1"""
                )
                actor = cursor.fetchone()
                if not actor:
                    raise RuntimeError(
                        "an active app user is required to attribute the historical import"
                    )
                actor_id = actor[0]

                # These are deliberately additive and remain compatible with
                # atelier.py's idempotent CREATE TABLE IF NOT EXISTS statements.
                for statement in (
                    "ALTER TABLE atelier_jobs ADD COLUMN IF NOT EXISTS source_system TEXT",
                    "ALTER TABLE atelier_jobs ADD COLUMN IF NOT EXISTS source_identity TEXT",
                    "ALTER TABLE atelier_jobs ADD COLUMN IF NOT EXISTS source_row INTEGER",
                    "ALTER TABLE atelier_jobs ADD COLUMN IF NOT EXISTS historical_customer_name TEXT",
                    "ALTER TABLE atelier_jobs ADD COLUMN IF NOT EXISTS historical_phone TEXT",
                    "ALTER TABLE atelier_jobs ADD COLUMN IF NOT EXISTS service_type TEXT",
                    "ALTER TABLE atelier_jobs ADD COLUMN IF NOT EXISTS date_in TIMESTAMPTZ",
                    "ALTER TABLE atelier_jobs ADD COLUMN IF NOT EXISTS collected_at TIMESTAMPTZ",
                    "ALTER TABLE atelier_jobs ADD COLUMN IF NOT EXISTS legacy_notes TEXT",
                ):
                    cursor.execute(statement)
                cursor.execute(
                    """CREATE UNIQUE INDEX IF NOT EXISTS
                       uq_atelier_jobs_source_identity
                       ON atelier_jobs(source_system,source_identity)
                       WHERE source_identity IS NOT NULL"""
                )
                cursor.execute(
                    """CREATE TABLE IF NOT EXISTS atelier_history_imports (
                         source_identity TEXT PRIMARY KEY,
                         source_system TEXT NOT NULL,
                         source_row INTEGER NOT NULL,
                         source_item INTEGER NOT NULL,
                         source_sha256 TEXT NOT NULL,
                         job_id BIGINT NOT NULL UNIQUE
                           REFERENCES atelier_jobs(id) ON DELETE CASCADE,
                         imported_at TIMESTAMPTZ NOT NULL DEFAULT now()
                       )"""
                )

                cursor.execute(
                    "SELECT user_id,name FROM app_users WHERE status='active'"
                )
                users = {
                    str(name).strip().casefold(): user_id
                    for user_id, name in cursor.fetchall()
                    if name
                }

                for job in result.jobs:
                    cursor.execute(
                        "SELECT job_id FROM atelier_history_imports "
                        "WHERE source_identity=%s",
                        (job.source_identity,),
                    )
                    if cursor.fetchone():
                        existing += 1
                        continue

                    customer = None
                    if job.phone:
                        national = job.phone[4:]
                        candidates = [national, "0" + national, "254" + national]
                        cursor.execute(
                            """SELECT customer_id,store_id FROM all_customers
                               WHERE regexp_replace(COALESCE(phone,''),'[^0-9]','','g')
                                     = ANY(%s)
                               ORDER BY CASE WHEN store_id='vivofashiongroup'
                                             THEN 0 ELSE 1 END,
                                        customer_id,store_id LIMIT 1""",
                            (candidates,),
                        )
                        customer = cursor.fetchone()
                    if customer:
                        customer_id, customer_store_id = customer
                    else:
                        # atelier_jobs intentionally has no customer FK so old
                        # anonymous sheet rows can remain visible and measurable
                        # as data-health gaps, without creating fake customers.
                        customer_id = f"history:{job.source_identity[:24]}"
                        customer_store_id = "historical-unresolved"

                    assigned_to = (
                        users.get(job.staff_name.casefold()) if job.staff_name else None
                    )
                    garment_type = job.item_description or "Garment"
                    alteration_notes = job.alteration_detail
                    cursor.execute(
                        """INSERT INTO atelier_jobs (
                             claim_number,customer_id,customer_store_id,location_id,
                             garment_index,garment_type,sku,product_name,
                             alteration_notes,promised_at,status,service_charge,
                             amount_paid,assigned_to,created_by,created_at,updated_at,
                             completed_at,source_system,source_identity,source_row,
                             historical_customer_name,historical_phone,service_type,
                             date_in,collected_at,legacy_notes
                           ) VALUES (
                             'VJ-' || %s || '-' ||
                               lpad(nextval('atelier_claim_seq')::text,5,'0'),
                             %s,%s,%s,%s,%s,%s,%s,%s,%s,%s,0,0,%s,%s,%s,%s,
                             %s,%s,%s,%s,%s,%s,%s,%s,%s,%s
                           )
                           ON CONFLICT (source_system,source_identity)
                             WHERE source_identity IS NOT NULL DO NOTHING
                           RETURNING id""",
                        (
                            job.date_in.year,
                            str(customer_id),
                            str(customer_store_id),
                            location_id,
                            job.source_item,
                            garment_type,
                            job.sku,
                            job.item_description,
                            alteration_notes,
                            job.promised_date,
                            db_status[job.status],
                            assigned_to,
                            actor_id,
                            job.date_in,
                            job.date_in,
                            job.date_ready or job.collected_at,
                            SOURCE_SYSTEM,
                            job.source_identity,
                            job.source_row,
                            job.customer_name,
                            job.phone,
                            job.service_type,
                            job.date_in,
                            job.collected_at,
                            job.notes,
                        ),
                    )
                    created = cursor.fetchone()
                    if not created:
                        existing += 1
                        continue
                    job_id = created[0]
                    cursor.execute(
                        """INSERT INTO atelier_status_history
                           (job_id,from_status,to_status,note,changed_by,changed_at)
                           VALUES (%s,NULL,%s,%s,%s,%s)""",
                        (
                            job_id,
                            db_status[job.status],
                            "Historical Tracker import",
                            actor_id,
                            job.collected_at or job.date_ready or job.date_in,
                        ),
                    )
                    cursor.execute(
                        """INSERT INTO atelier_history_imports
                           (source_identity,source_system,source_row,source_item,
                            source_sha256,job_id)
                           VALUES (%s,%s,%s,%s,%s,%s)""",
                        (
                            job.source_identity,
                            SOURCE_SYSTEM,
                            job.source_row,
                            job.source_item,
                            result.source_sha256,
                            job_id,
                        ),
                    )
                    inserted += 1
    finally:
        connection.close()
    return inserted, existing


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", nargs="?", default=DEFAULT_WORKBOOK)
    parser.add_argument(
        "--commit",
        action="store_true",
        help="write importable rows; without this flag the command is read-only",
    )
    parser.add_argument("--database-url", default=None, help=argparse.SUPPRESS)
    return parser


def main(argv: Iterable[str] | None = None) -> int:
    args = build_parser().parse_args(list(argv) if argv is not None else None)
    try:
        result = parse_workbook(args.workbook)
        inserted = existing = 0
        if args.commit:
            database_url = args.database_url or os.environ.get("DATABASE_URL")
            if not database_url:
                raise RuntimeError("DATABASE_URL is required with --commit")
            inserted, existing = commit_jobs(result, database_url)
        print(
            json.dumps(
                result.report(
                    committed=args.commit, inserted=inserted, existing=existing
                ),
                indent=2,
                sort_keys=True,
                default=str,
            )
        )
        return 0
    except Exception as exc:
        print(json.dumps({"error": str(exc), "mode": "commit" if args.commit else "dry-run"}))
        return 1


if __name__ == "__main__":
    sys.exit(main())