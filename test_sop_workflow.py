import asyncio
import io
import os
import re
import unicodedata
import unittest
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from fastapi import HTTPException
from pypdf import PdfReader

import api_pg


def _request(user):
    return SimpleNamespace(state=SimpleNamespace(user=user))


class SopWorkflowRulesTests(unittest.TestCase):
    @staticmethod
    def _docx_bytes():
        from docx import Document

        output = io.BytesIO()
        document = Document()
        document.add_heading("Store Opening", level=1)
        paragraph = document.add_paragraph()
        paragraph.add_run("Open the store ").bold = True
        paragraph.add_run("safely.")
        document.add_paragraph("Check the alarm", style="List Bullet")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Owner"
        table.cell(0, 1).text = "Task"
        table.cell(1, 0).text = "Manager"
        table.cell(1, 1).text = "Unlock doors"
        document.save(output)
        return output.getvalue()

    def test_fixed_stage_and_department_taxonomy(self):
        self.assertEqual(
            [stage["name"] for stage in api_pg.SOP_STAGES],
            [
                "01. SOP Submission",
                "02. SOPs Under Review",
                "03. Awaiting Approval",
                "04. Approved SOPs (Master Repository)",
                "05. Obsolete SOPs",
            ],
        )
        self.assertEqual(len(api_pg.SOP_DEPARTMENTS), 11)
        self.assertEqual(
            api_pg.SOP_DEPARTMENTS[-1]["name"], "Warehouse & Logistics")

    def test_regular_users_only_access_submission_and_approved(self):
        user = {"role": "retail", "email": "staff@vivofashiongroup.com"}
        self.assertTrue(api_pg._sop_can_access_stage(user, 1))
        self.assertFalse(api_pg._sop_can_access_stage(user, 2))
        self.assertFalse(api_pg._sop_can_access_stage(user, 5))
        self.assertTrue(api_pg._sop_can_access_stage(user, 3))
        self.assertFalse(api_pg._sop_can_access_stage(user, 4))

    def test_named_reviewers_and_approvers(self):
        frankie = {
            "role": "retail",
            "email": "FRANCKIE@VIVOFASHIONGROUP.COM",
        }
        stephen = {
            "role": "leadership",
            "email": "stephen@vivofashiongroup.com",
        }
        self.assertTrue(api_pg._sop_can_access_stage(frankie, 2))
        self.assertTrue(api_pg._sop_can_access_stage(frankie, 5))
        self.assertTrue(api_pg._sop_can_review(frankie))
        self.assertTrue(api_pg._sop_can_edit_stage(frankie, 2))
        self.assertFalse(api_pg._sop_can_edit_stage(frankie, 5))
        self.assertFalse(api_pg._sop_can_approve(frankie))
        self.assertTrue(api_pg._sop_can_edit_stage(stephen, 5))
        self.assertTrue(api_pg._sop_can_approve(stephen))

    @patch.object(api_pg, "_ensure_sop_tables")
    @patch.object(api_pg, "_users_exec")
    def test_hidden_stage_is_rejected_before_query(self, users_exec, ensure):
        with self.assertRaises(HTTPException) as raised:
            api_pg.sops_files(
                _request({
                    "role": "marketing",
                    "email": "staff@vivofashiongroup.com",
                }),
                stage=2,
                department="marketing",
            )
        self.assertEqual(raised.exception.status_code, 403)
        ensure.assert_not_called()
        users_exec.assert_not_called()

    @patch.object(api_pg, "_log_activity")
    @patch.object(api_pg, "_ensure_sop_tables")
    @patch.object(api_pg, "_users_exec")
    def test_frankie_can_atomically_review_stage_one(
        self, users_exec, ensure, log_activity
    ):
        users_exec.return_value = [{
            "id": 17,
            "filename": "Store opening SOP.pdf",
            "department": "retail-operations",
            "stage": 2,
        }]
        result = api_pg.sops_review(
            17,
            _request({
                "role": "retail",
                "email": "franckie@vivofashiongroup.com",
            }),
        )
        self.assertTrue(result["ok"])
        self.assertIn("WHERE id=%s AND stage=1", users_exec.call_args.args[0])
        log_activity.assert_called_once()

    def test_frankie_cannot_approve(self):
        with self.assertRaises(HTTPException) as raised:
            api_pg.sops_approve(
                17,
                _request({
                    "role": "retail",
                    "email": "franckie@vivofashiongroup.com",
                }),
            )
        self.assertEqual(raised.exception.status_code, 403)

    @patch.object(api_pg, "_log_activity")
    @patch.object(api_pg, "_ensure_sop_tables")
    @patch.object(api_pg, "_sop_user_grants", return_value={"finance"})
    @patch.object(api_pg, "_users_exec")
    def test_submitter_delete_cannot_cross_review_transition(
        self, users_exec, grants, ensure, log_activity
    ):
        users_exec.side_effect = [
            [{
                "id": 17,
                "stage": 1,
                "department": "finance",
                "filename": "Finance SOP.pdf",
            }],
            [],
        ]
        with self.assertRaises(HTTPException) as raised:
            api_pg.sops_delete(
                17,
                _request({
                    "user_id": "submitter-1",
                    "role": "finance",
                    "email": "submitter@vivofashiongroup.com",
                }),
            )
        self.assertEqual(raised.exception.status_code, 409)
        delete_sql = users_exec.call_args_list[1].args[0]
        self.assertIn("stage=1", delete_sql)
        log_activity.assert_not_called()

    @patch.object(api_pg, "_log_activity")
    @patch.object(api_pg, "_ensure_sop_tables")
    @patch.object(api_pg, "_users_exec")
    def test_frankie_can_atomically_retire_approved_sop(
        self, users_exec, ensure, log_activity
    ):
        users_exec.return_value = [{
            "id": 23,
            "filename": "Store opening SOP.pdf",
            "department": "retail-operations",
            "stage": 4,
        }]
        result = api_pg.sops_obsolete(
            23,
            _request({
                "role": "retail",
                "email": "franckie@vivofashiongroup.com",
            }),
        )
        self.assertTrue(result["ok"])
        self.assertIn("WHERE id=%s AND stage=3", users_exec.call_args.args[0])
        self.assertIn("SELECT id, 3, 4", users_exec.call_args.args[0])
        log_activity.assert_called_once()

    def test_editor_html_is_sanitized(self):
        value = (
            '<h1 onclick="bad()">Title</h1><script>alert(1)</script>'
            '<p style="position:fixed">Safe <strong>text</strong></p>'
        )
        cleaned = api_pg._sop_sanitize_html(value)
        self.assertEqual(
            cleaned,
            "<h1>Title</h1>alert(1)<p>Safe <strong>text</strong></p>",
        )
        self.assertNotIn("onclick", cleaned)
        self.assertNotIn("<script", cleaned)

    @patch.object(api_pg, "_ensure_sop_tables")
    @patch.object(api_pg, "_users_exec")
    def test_editor_replaces_legacy_placeholder_with_original_docx(
        self, users_exec, ensure
    ):
        users_exec.return_value = [{
            "id": 7,
            "stage": 2,
            "department": "finance",
            "filename": "Odoo_Remediation_Plan.docx",
            "data": self._docx_bytes(),
            "content_type": (
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            "original_filename": "Odoo_Remediation_Plan.docx",
            "original_data": self._docx_bytes(),
            "original_content_type": (
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            "editor_html": (
                "<h1>Odoo_Remediation_Plan.docx</h1>"
                "<p>This SOP is ready for editing in the dashboard. "
                "The original uploaded file is retained for traceability.</p>"
            ),
            "editor_revision": 0,
            "edited_by_email": None,
            "edited_at": None,
        }]
        result = api_pg.sops_editor(
            7,
            _request({
                "role": "retail",
                "email": "franckie@vivofashiongroup.com",
            }),
        )
        self.assertIn("<h1>Store Opening</h1>", result["html"])
        self.assertNotIn("ready for editing", result["html"])
        self.assertIn("original_filename", users_exec.call_args.args[0])

    def test_docx_contents_are_imported_into_editable_html(self):
        imported = api_pg._sop_initial_editor_html(
            "Store Opening.docx",
            self._docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertIn("<h1>Store Opening</h1>", imported)
        self.assertIn("<strong>Open the store </strong>", imported)
        self.assertIn("<ul><li>Check the alarm</li></ul>", imported)
        self.assertIn("<table>", imported)
        self.assertIn("<th><p>Owner</p></th>", imported)
        self.assertIn("<td><p>Unlock doors</p></td>", imported)
        self.assertNotIn("ready for editing", imported)

    def test_empty_word_note_parts_and_font_metadata_are_not_rejected(self):
        from docx import Document
        from docx.shared import Pt, RGBColor

        document = Document()
        paragraph = document.add_paragraph()
        run = paragraph.add_run("Normal styled procedure")
        run.font.name = "Arial"
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(31, 41, 55)
        output = io.BytesIO()
        document.save(output)

        imported = api_pg._sop_initial_editor_html(
            "Styled procedure.docx",
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertIn("Normal styled procedure", imported)

    def test_approved_pdf_contains_the_final_editor_text(self):
        pdf = api_pg._sop_html_to_pdf(
            "<h1>Store Opening</h1><p>Final approved procedure — café 中文</p>"
            "<ul><li>Check the alarm</li></ul>"
        )
        self.assertTrue(pdf.startswith(b"%PDF"))
        text = "\n".join(
            page.extract_text() or ""
            for page in PdfReader(io.BytesIO(pdf)).pages
        )
        self.assertIn("Store Opening", text)
        self.assertIn("Final approved procedure", text)
        self.assertIn("café", text)
        self.assertIn("中文", text)
        self.assertIn("Check the alarm", text)

    def test_approved_pdf_round_trips_shaped_multilingual_text_and_emoji(self):
        samples = (
            "Latin café",
            "مرحبا بالعالم",
            "שלום עולם",
            "हिन्दी भाषा",
            "ไทย ภาษาไทย",
            "中文 日本語 한국어",
            "Staff greeting 😀",
        )
        for sample in samples:
            with self.subTest(sample=sample):
                pdf = api_pg._sop_html_to_pdf(f"<p>{sample}</p>")
                text = unicodedata.normalize(
                    "NFC",
                    "\n".join(
                        page.extract_text() or ""
                        for page in PdfReader(io.BytesIO(pdf)).pages
                    ),
                )
                self.assertIn(unicodedata.normalize("NFC", sample), text)

        combining_pdf = api_pg._sop_html_to_pdf("<p>Combining: e\u0301</p>")
        combining_text = unicodedata.normalize(
            "NFC",
            "\n".join(
                page.extract_text() or ""
                for page in PdfReader(io.BytesIO(combining_pdf)).pages
            ),
        )
        self.assertIn("Combining: é", combining_text)

        combined_pdf = api_pg._sop_html_to_pdf(
            "".join(f"<p>{sample}</p>" for sample in samples))
        reader = PdfReader(io.BytesIO(combined_pdf))
        embedded_fonts = set()
        embedded_font_files = set()
        for page in reader.pages:
            fonts = page["/Resources"]["/Font"].get_object()
            for font_ref in fonts.values():
                font = font_ref.get_object()
                base_font = str(font.get("/BaseFont", ""))
                embedded_fonts.add(base_font)
                descendants = font.get("/DescendantFonts")
                if not descendants:
                    continue
                descendant = descendants[0].get_object()
                descriptor = descendant["/FontDescriptor"].get_object()
                if any(key in descriptor
                       for key in ("/FontFile", "/FontFile2", "/FontFile3")):
                    embedded_font_files.add(
                        str(descriptor.get("/FontFamily", "")))
        for family in (
            "SopNoto", "SopArabic", "SopHebrew", "SopDevanagari",
            "SopThai", "SopCJK", "SopEmoji",
        ):
            with self.subTest(family=family):
                self.assertTrue(
                    any(family in font for font in embedded_fonts),
                    f"{family} was not embedded in the approved PDF",
                )
                self.assertIn(
                    family,
                    embedded_font_files,
                    f"{family} has no embedded FontFile stream",
                )

    def test_approved_pdf_contains_shaped_complex_script_glyph_runs(self):
        def first_text_array(text):
            pdf = api_pg._sop_html_to_pdf(f"<p>{text}</p>")
            content = PdfReader(io.BytesIO(pdf)).pages[0].get_contents(
            ).get_data()
            match = re.search(rb"\[(.*?)\]\s*TJ", content, re.DOTALL)
            self.assertIsNotNone(match)
            return match.group(1)

        def glyph_codes(text_array):
            return [
                value for value in re.findall(rb"<([0-9a-fA-F]*)>",
                                               text_array) if value
            ]

        arabic_shaped = glyph_codes(first_text_array("سلام"))
        arabic_unjoined = glyph_codes(first_text_array("س\u200cل\u200cا\u200cم"))
        self.assertEqual(len(arabic_shaped), 4)
        self.assertEqual(len(arabic_unjoined), 4)
        self.assertNotEqual(
            arabic_shaped,
            arabic_unjoined,
            "Arabic contextual forms were not substituted",
        )

        devanagari_conjunct = glyph_codes(first_text_array("क्ष"))
        devanagari_broken = glyph_codes(first_text_array("क्\u200cष"))
        self.assertEqual(
            len(devanagari_conjunct),
            1,
            "Devanagari conjunct was not shaped into a single glyph",
        )
        self.assertGreaterEqual(len(devanagari_broken), 3)

        thai_mark_run = first_text_array("ป่")
        positioning = [
            float(value)
            for value in re.findall(rb">(-?\d+(?:\.\d+)?)<", thai_mark_run)
        ]
        self.assertTrue(
            any(abs(value) > 0.01 for value in positioning),
            "Thai combining mark had no shaped positioning adjustment",
        )

    def test_approved_pdf_fonts_are_packaged_with_the_application(self):
        from fontTools.ttLib import TTFont

        for _family, filename, _style, _weight in api_pg._SOP_PDF_FONTS:
            with self.subTest(filename=filename):
                path = os.path.join(api_pg._SOP_FONT_DIR, filename)
                self.assertTrue(os.path.isfile(path), path)
                self.assertGreater(os.path.getsize(path), 1000)
                font = TTFont(path, lazy=True)
                try:
                    if filename in {
                        "NotoSansArabic-Regular.ttf",
                        "NotoSansDevanagari-Regular.ttf",
                        "NotoSansThai-Regular.ttf",
                    }:
                        self.assertIn("GSUB", font)
                        self.assertIn("GPOS", font)
                finally:
                    font.close()

    def test_approved_pdf_uses_visual_rtl_order_for_arabic_and_hebrew(self):
        for text in ("مرحبا بالعالم", "שלום עולם"):
            with self.subTest(text=text):
                pdf = api_pg._sop_html_to_pdf(f"<p>{text}</p>")
                layout = PdfReader(io.BytesIO(pdf)).pages[0].extract_text(
                    extraction_mode="layout")
                first_content_line = next(
                    line.strip() for line in layout.splitlines() if line.strip())
                # Layout extraction walks glyphs left-to-right. A correctly
                # bidi-laid-out RTL line therefore appears in reverse logical
                # codepoint order here.
                self.assertEqual(first_content_line, text[::-1])

    def test_approved_pdf_preserves_table_rows_and_columns(self):
        pdf = api_pg._sop_html_to_pdf(
            "<h1>Opening Roles</h1>"
            "<table><tbody>"
            "<tr><th><p>Owner</p></th><th><p>Action</p></th></tr>"
            "<tr><td><p>Manager</p></td><td><p>Unlock doors</p></td></tr>"
            "</tbody></table>"
        )
        layout = "\n".join(
            page.extract_text(extraction_mode="layout") or ""
            for page in PdfReader(io.BytesIO(pdf)).pages
        )
        lines = [" ".join(line.split()) for line in layout.splitlines()]
        self.assertIn("Owner Action", lines)
        self.assertIn("Manager Unlock doors", lines)

    def test_docx_numbered_steps_remain_numbered_in_approved_pdf(self):
        from docx import Document

        doc = Document()
        doc.add_heading("Closing Procedure", level=1)
        doc.add_paragraph("Lock the stock room", style="List Number")
        doc.add_paragraph("Set the alarm", style="List Number")
        raw = io.BytesIO()
        doc.save(raw)

        imported = api_pg._sop_initial_editor_html(
            "Closing.docx",
            raw.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertIn(
            "<ol><li>Lock the stock room</li><li>Set the alarm</li></ol>",
            imported,
        )
        pdf = api_pg._sop_html_to_pdf(imported)
        text = "\n".join(
            page.extract_text(extraction_mode="layout") or ""
            for page in PdfReader(io.BytesIO(pdf)).pages
        )
        self.assertIn("1. Lock the stock room", text)
        self.assertIn("2. Set the alarm", text)
        self.assertNotIn("• Lock the stock room", text)

    def test_approval_rejects_an_unavailable_glyph_without_silent_loss(self):
        with self.assertRaises(ValueError) as raised:
            api_pg._sop_html_to_pdf("<p>Unavailable \U0010ffff</p>")
        self.assertIn("cannot safely preserve", str(raised.exception))
        self.assertIn("U+10FFFF", str(raised.exception))

    @patch.object(api_pg, "_log_activity")
    @patch.object(api_pg, "_ensure_sop_tables")
    @patch.object(api_pg, "_users_exec")
    @patch.object(
        api_pg, "_sop_html_to_pdf",
        side_effect=RuntimeError("renderer unavailable"))
    def test_editor_approval_returns_400_before_update_when_rendering_fails(
        self, render_pdf, users_exec, ensure, log_activity
    ):
        users_exec.return_value = [{
            "stage": 5,
            "department": "finance",
            "filename": "Policy.docx",
            "original_filename": "Policy.docx",
            "original_data": self._docx_bytes(),
        }]
        with self.assertRaises(HTTPException) as raised:
            api_pg.sops_editor_update(
                31,
                _request({
                    "role": "leadership",
                    "email": "stephen@vivofashiongroup.com",
                }),
                {
                    "html": "<h1>Policy</h1>",
                    "revision": 4,
                    "transition": "approve",
                },
            )
        self.assertEqual(raised.exception.status_code, 400)
        self.assertIn("could not be rendered safely", raised.exception.detail)
        self.assertEqual(users_exec.call_count, 1)
        self.assertTrue(users_exec.call_args.args[0].startswith("SELECT stage"))
        log_activity.assert_not_called()

    @patch.object(api_pg, "_log_activity")
    @patch.object(api_pg, "_ensure_sop_tables")
    @patch.object(api_pg, "_users_exec")
    @patch.object(
        api_pg, "_sop_html_to_pdf",
        side_effect=RuntimeError("renderer unavailable"))
    def test_direct_approval_returns_400_before_update_when_rendering_fails(
        self, render_pdf, users_exec, ensure, log_activity
    ):
        users_exec.return_value = [{
            "stage": 5,
            "filename": "Policy.docx",
            "editor_html": "<h1>Policy</h1>",
            "editor_revision": 4,
            "original_filename": "Policy.docx",
            "original_data": self._docx_bytes(),
        }]
        with self.assertRaises(HTTPException) as raised:
            api_pg.sops_approve(
                31,
                _request({
                    "role": "leadership",
                    "email": "stephen@vivofashiongroup.com",
                }),
            )
        self.assertEqual(raised.exception.status_code, 400)
        self.assertIn("could not be rendered safely", raised.exception.detail)
        self.assertEqual(users_exec.call_count, 1)
        self.assertTrue(users_exec.call_args.args[0].startswith("SELECT stage"))
        log_activity.assert_not_called()

    @patch.object(api_pg, "_users_exec")
    @patch.object(api_pg, "_ensure_sop_tables")
    @patch.object(api_pg, "_sop_can_upload", return_value=True)
    def test_malformed_docx_returns_a_clear_upload_error(
        self, can_upload, ensure, users_exec
    ):
        upload = SimpleNamespace(
            filename="Broken.docx",
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            read=AsyncMock(return_value=b"PKnot-a-real-word-file"),
        )
        with self.assertRaises(HTTPException) as raised:
            asyncio.run(api_pg.sops_upload(
                _request({
                    "user_id": "submitter-1",
                    "role": "finance",
                    "email": "submitter@vivofashiongroup.com",
                }),
                department="finance",
                file=upload,
            ))
        self.assertEqual(raised.exception.status_code, 400)
        self.assertIn("Could not read this Word document", raised.exception.detail)
        users_exec.assert_not_called()

    @patch.object(api_pg, "_ensure_sop_tables")
    @patch.object(api_pg, "_sop_can_upload", return_value=True)
    def test_upload_rejects_non_word_documents(self, can_upload, ensure):
        upload = SimpleNamespace(
            filename="Policy.pdf",
            content_type="application/pdf",
            read=AsyncMock(return_value=b"%PDF"),
        )
        with self.assertRaises(HTTPException) as raised:
            asyncio.run(api_pg.sops_upload(
                _request({
                    "user_id": "submitter-1",
                    "role": "finance",
                    "email": "submitter@vivofashiongroup.com",
                }),
                department="finance",
                file=upload,
            ))
        self.assertEqual(raised.exception.status_code, 400)
        self.assertIn("Only modern Word documents", raised.exception.detail)
        upload.read.assert_not_awaited()

    def test_docx_with_header_is_rejected_instead_of_silently_dropping_it(self):
        from docx import Document

        doc = Document()
        doc.add_paragraph("Body procedure")
        doc.sections[0].header.paragraphs[0].text = "Mandatory header"
        raw = io.BytesIO()
        doc.save(raw)
        with self.assertRaises(ValueError) as raised:
            api_pg._sop_initial_editor_html("Header.docx", raw.getvalue())
        self.assertIn("headers or footers", str(raised.exception))

    def test_docx_with_merged_table_is_rejected(self):
        from docx import Document

        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged instruction"
        raw = io.BytesIO()
        doc.save(raw)
        with self.assertRaises(ValueError) as raised:
            api_pg._sop_initial_editor_html("Merged.docx", raw.getvalue())
        self.assertIn("merged table cells", str(raised.exception))

    def test_docx_with_nested_numbered_list_is_rejected(self):
        from docx import Document

        doc = Document()
        doc.add_paragraph("Main step", style="List Number")
        doc.add_paragraph("Nested step", style="List Number 2")
        raw = io.BytesIO()
        doc.save(raw)
        with self.assertRaises(ValueError) as raised:
            api_pg._sop_initial_editor_html("Nested.docx", raw.getvalue())
        self.assertIn("multilevel lists", str(raised.exception))

    def test_legacy_doc_is_rejected_instead_of_flattened(self):
        with self.assertRaises(ValueError) as raised:
            api_pg._sop_initial_editor_html(
                "Legacy.doc",
                b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1legacy",
            )
        self.assertIn("Only modern Word documents", str(raised.exception))

    @patch.object(api_pg, "_log_activity")
    @patch.object(api_pg, "_ensure_sop_tables")
    @patch.object(api_pg, "_users_exec")
    def test_frankie_edit_moves_under_review_to_awaiting_approval(
        self, users_exec, ensure, log_activity
    ):
        users_exec.side_effect = [
            [{"stage": 2, "department": "finance", "filename": "Policy.docx"}],
            [{
                "id": 31,
                "stage": 5,
                "filename": "Policy.docx",
                "department": "finance",
                "editor_revision": 4,
                "edited_by_email": "franckie@vivofashiongroup.com",
                "edited_at": None,
            }],
        ]
        result = api_pg.sops_editor_update(
            31,
            _request({
                "role": "retail",
                "email": "franckie@vivofashiongroup.com",
            }),
            {"html": "<h1>Edited</h1>", "revision": 3,
             "transition": "awaiting_approval"},
        )
        self.assertTrue(result["ok"])
        transition_sql = users_exec.call_args_list[1].args[0]
        self.assertIn("SET stage=%s", transition_sql)
        self.assertIn("SELECT id, %s, %s", transition_sql)
        self.assertIn("INSERT INTO sop_file_revisions", transition_sql)
        self.assertEqual(users_exec.call_args_list[1].args[1][0], 5)
        log_activity.assert_called_once()

    @patch.object(api_pg, "_log_activity")
    @patch.object(api_pg, "_ensure_sop_tables")
    @patch.object(api_pg, "_users_exec")
    def test_stephen_edit_approves_awaiting_sop(
        self, users_exec, ensure, log_activity
    ):
        users_exec.side_effect = [
            [{
                "stage": 5,
                "department": "finance",
                "filename": "Policy.docx",
                "original_filename": "Policy.docx",
                "original_data": self._docx_bytes(),
            }],
            [{
                "id": 31,
                "stage": 3,
                "filename": "Policy.pdf",
                "department": "finance",
                "editor_revision": 5,
                "edited_by_email": "stephen@vivofashiongroup.com",
                "edited_at": None,
            }],
        ]
        result = api_pg.sops_editor_update(
            31,
            _request({
                "role": "leadership",
                "email": "stephen@vivofashiongroup.com",
            }),
            {"html": "<h1>Approved</h1>", "revision": 4,
             "transition": "approve"},
        )
        self.assertEqual(result["file"]["stage"], 3)
        approval_params = users_exec.call_args_list[1].args[1]
        self.assertEqual(approval_params[0], 3)
        self.assertEqual(approval_params[1], "Policy.pdf")
        self.assertEqual(approval_params[5], "application/pdf")
        self.assertTrue(bytes(approval_params[4].adapted).startswith(b"%PDF"))
        self.assertIn(
            "INSERT INTO sop_file_revisions",
            users_exec.call_args_list[1].args[0],
        )
        log_activity.assert_called_once()

    @patch.object(api_pg, "_log_activity")
    @patch.object(api_pg, "_ensure_sop_tables")
    @patch.object(api_pg, "_users_exec")
    def test_stale_editor_revision_is_rejected(
        self, users_exec, ensure, log_activity
    ):
        users_exec.side_effect = [
            [{"stage": 2, "department": "finance", "filename": "Policy.docx"}],
            [],
            [{"stage": 2, "editor_revision": 8}],
        ]
        with self.assertRaises(HTTPException) as raised:
            api_pg.sops_editor_update(
                31,
                _request({
                    "role": "retail",
                    "email": "franckie@vivofashiongroup.com",
                }),
                {"html": "<p>Stale edit</p>", "revision": 7,
                 "transition": "save"},
            )
        self.assertEqual(raised.exception.status_code, 409)
        log_activity.assert_not_called()


if __name__ == "__main__":
    unittest.main()
