#!/usr/bin/env python3
"""One-off / re-runnable loader for the sales-targets store.

Populates `targets_monthly` with two layers:

  • scope='region', source='budget'  — the leadership "Revenues Budget 2026"
    workbook, summarised to the four annual-targets buckets the Targets page
    expects: "Kenya - Retail", "Kenya - Online", "Uganda", "Rwanda". These
    reconcile to the workbook's own Quarterly Summary sheet (Kenya-Online is
    the "Gross SZ" row, per that summary).

  • scope='store',  source='manual'  — the per-store monthly targets the
    leadership team types in (June 2026 to start; more months added over time).
    Names are mapped to the canonical all_sales.pos_location_name values so the
    monthly tracker can join actuals.

Idempotent: re-running upserts on (scope, name, month, source).

Run:  python3 load_targets.py
"""
import os
import sys
from datetime import date

import psycopg2

BUDGET_XLSX = "attached_assets/Revenues_Budget_2026_-_Team_1781772723098.xlsx"
YEAR = 2026

MONTHS = ["January", "February", "March", "April", "May", "June",
          "July", "August", "September", "October", "November", "December"]

# Typed per-store June 2026 targets → canonical all_sales.pos_location_name.
# "Greenwood" is the Meru-Greenwood store (workbook row "Meru-Greenwood"),
# which posts to all_sales as "Vivo Meru".
JUNE_2026_STORES = {
    "Vivo Sarit": (8_600_000, "Kenya"),
    "Vivo Junction": (8_600_000, "Kenya"),
    "Vivo Mama Ngina St": (6_000_000, "Kenya"),
    "Vivo Moi Avenue": (5_500_000, "Kenya"),
    "Vivo Yaya": (5_000_000, "Kenya"),
    "Vivo Garden City": (3_900_000, "Kenya"),
    "Vivo TRM": (3_900_000, "Kenya"),
    "Vivo Village Market": (5_200_000, "Kenya"),
    "Vivo Two Rivers": (3_500_000, "Kenya"),
    "Vivo Capital Centre": (3_500_000, "Kenya"),
    "Vivo Galleria": (3_300_000, "Kenya"),
    "Vivo Hub": (3_000_000, "Kenya"),
    "Vivo Imaara": (3_700_000, "Kenya"),
    "Vivo T- Mall": (2_300_000, "Kenya"),
    "Vivo Kileleshwa": (1_600_000, "Kenya"),
    "Vivo Signature Mall": (2_000_000, "Kenya"),
    "Vivo Greenspan": (2_000_000, "Kenya"),
    "Vivo Nakuru": (3_400_000, "Kenya"),
    "Vivo Eldoret": (3_100_000, "Kenya"),
    "Vivo Kisumu": (2_900_000, "Kenya"),
    "Vivo City Mall": (3_100_000, "Kenya"),
    "Vivo MSA Digo Road": (2_000_000, "Kenya"),
    "Vivo Meru": (1_200_000, "Kenya"),
    "The Oasis Mall": (2_900_000, "Uganda"),
    "Vivo Kigali Heights": (4_400_000, "Rwanda"),
    "Safari Sarit": (1_600_000, "Kenya"),
    "Zoya Sarit": (900_000, "Kenya"),
    "Vivo Runda": (3_100_000, "Kenya"),
    "Vivo Acacia": (5_000_000, "Uganda"),
}

# Typed per-store July 2026 targets → canonical all_sales.pos_location_name.
# "Greenwood" is the Meru-Greenwood store (posts as "Vivo Meru"). The workbook
# line "Sarit Safari/Zoya" (3,190,000) is split evenly across the two Sarit
# sub-stores (Safari Sarit / Zoya Sarit) per leadership instruction, and the
# Online channel is tracked as the "Online - Shop Zetu" pos_location_name.
JULY_2026_STORES = {
    "Vivo Sarit": (9_035_000, "Kenya"),
    "Vivo Junction": (9_035_000, "Kenya"),
    "Vivo Mama Ngina St": (7_440_000, "Kenya"),
    "Vivo Moi Avenue": (6_910_000, "Kenya"),
    "Vivo Village Market": (6_110_000, "Kenya"),
    "Vivo Yaya": (6_110_000, "Kenya"),
    "Vivo Garden City": (4_785_000, "Kenya"),
    "Vivo Galleria": (4_465_000, "Kenya"),
    "Vivo Imaara": (4_040_000, "Kenya"),
    "Vivo Nakuru": (4_040_000, "Kenya"),
    "Vivo TRM": (3_935_000, "Kenya"),
    "Vivo Two Rivers": (3_720_000, "Kenya"),
    "Vivo Capital Centre": (3_720_000, "Kenya"),
    "Vivo Eldoret": (3_615_000, "Kenya"),
    "Vivo City Mall": (3_510_000, "Kenya"),
    "Vivo Hub": (3_510_000, "Kenya"),
    "Vivo Kisumu": (3_510_000, "Kenya"),
    "Vivo Runda": (3_405_000, "Kenya"),
    "Safari Sarit": (1_595_000, "Kenya"),
    "Zoya Sarit": (1_595_000, "Kenya"),
    "Vivo MSA Digo Road": (2_925_000, "Kenya"),
    "Vivo Signature Mall": (2_660_000, "Kenya"),
    "Vivo T- Mall": (2_555_000, "Kenya"),
    "Vivo Kileleshwa": (2_445_000, "Kenya"),
    "Vivo Meru": (2_130_000, "Kenya"),
    "Vivo Greenspan": (2_130_000, "Kenya"),
    "Vivo Kigali Heights": (4_680_000, "Rwanda"),
    "The Oasis Mall": (3_550_000, "Uganda"),
    "Vivo Acacia": (6_100_000, "Uganda"),
    "Online - Shop Zetu": (10_000_000, "Online"),
}

# Per-store targets keyed by month number (add new months here over time).
STORE_TARGETS_BY_MONTH = {
    6: JUNE_2026_STORES,
    7: JULY_2026_STORES,
}


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def extract_region_budget():
    """Return {bucket: {month_int: target_kes}} for the 4 annual buckets."""
    import openpyxl
    wb = openpyxl.load_workbook(BUDGET_XLSX, data_only=True)

    def month_cols(ws):
        for row in ws.iter_rows(min_row=1, max_row=4, values_only=True):
            if row and any(c == "January" for c in row if isinstance(c, str)):
                idx = {}
                for i, c in enumerate(row):
                    if isinstance(c, str) and c.strip() in MONTHS:
                        idx[MONTHS.index(c.strip()) + 1] = i
                return idx
        raise RuntimeError("No month header row found")

    out = {"Kenya - Retail": {}, "Kenya - Online": {},
           "Uganda": {}, "Rwanda": {}}

    # Kenya retail = sum of every store row (skip the 'Total' row).
    ws = wb["Kenya"]
    mc = month_cols(ws)
    for q in range(1, 13):
        out["Kenya - Retail"][q] = 0.0
    for row in ws.iter_rows(values_only=True):
        label = row[1] if len(row) > 1 else None
        if not isinstance(label, str) or not label.strip():
            continue
        if label.strip().lower() in ("total", "projected monthly sales"):
            continue
        for q, ci in mc.items():
            v = _num(row[ci]) if ci < len(row) else None
            if v:
                out["Kenya - Retail"][q] += v

    # Kenya online = "Gross SZ" row (matches the Quarterly Summary).
    ws = wb["Online"]
    mc = month_cols(ws)
    for row in ws.iter_rows(values_only=True):
        label = row[1] if len(row) > 1 else None
        if isinstance(label, str) and label.strip().lower() == "gross sz":
            for q, ci in mc.items():
                out["Kenya - Online"][q] = _num(row[ci]) or 0.0

    # Uganda = Oasis + Acacia store rows.
    ws = wb["Uganda"]
    mc = month_cols(ws)
    for q in range(1, 13):
        out["Uganda"][q] = 0.0
    for row in ws.iter_rows(values_only=True):
        label = row[1] if len(row) > 1 else None
        if isinstance(label, str) and label.strip() in ("Oasis Mall", "Acacia Mall"):
            for q, ci in mc.items():
                v = _num(row[ci]) if ci < len(row) else None
                if v:
                    out["Uganda"][q] += v

    # Rwanda = Kigali Heights row.
    ws = wb["Rwanda"]
    mc = month_cols(ws)
    for row in ws.iter_rows(values_only=True):
        label = row[1] if len(row) > 1 else None
        if isinstance(label, str) and label.strip() == "Kigali Heights":
            for q, ci in mc.items():
                out["Rwanda"][q] = _num(row[ci]) or 0.0

    return out


DDL = """
CREATE TABLE IF NOT EXISTS targets_monthly (
    id          BIGSERIAL PRIMARY KEY,
    scope       TEXT NOT NULL,            -- 'region' | 'store'
    name        TEXT NOT NULL,            -- bucket name | canonical pos_location_name
    country     TEXT,                     -- region / country grouping
    month       DATE NOT NULL,            -- first day of the month
    target_kes  NUMERIC NOT NULL,
    source      TEXT NOT NULL,            -- 'budget' | 'manual'
    updated_at  TIMESTAMP NOT NULL DEFAULT now(),
    UNIQUE (scope, name, month, source)
);
"""


def main():
    dsn = os.environ.get("DATABASE_URL")
    if not dsn:
        print("DATABASE_URL not set", file=sys.stderr)
        sys.exit(1)
    region = extract_region_budget()

    conn = psycopg2.connect(dsn)
    conn.autocommit = False
    cur = conn.cursor()
    cur.execute(DDL)

    upsert = (
        "INSERT INTO targets_monthly (scope, name, country, month, target_kes, source) "
        "VALUES (%s,%s,%s,%s,%s,%s) "
        "ON CONFLICT (scope, name, month, source) DO UPDATE SET "
        "target_kes = EXCLUDED.target_kes, country = EXCLUDED.country, updated_at = now()"
    )

    region_country = {"Kenya - Retail": "Kenya", "Kenya - Online": "Online",
                      "Uganda": "Uganda", "Rwanda": "Rwanda"}
    n_region = 0
    for bucket, months in region.items():
        for m, val in months.items():
            cur.execute(upsert, ("region", bucket, region_country[bucket],
                                 date(YEAR, m, 1), round(val), "budget"))
            n_region += 1

    n_store = 0
    for month_num, stores in STORE_TARGETS_BY_MONTH.items():
        for store, (val, country) in stores.items():
            cur.execute(upsert, ("store", store, country,
                                 date(YEAR, month_num, 1), val, "manual"))
            n_store += 1

    conn.commit()

    # Report
    cur.execute("SELECT name, target_kes FROM targets_monthly "
                "WHERE scope='region' AND source='budget' "
                "AND EXTRACT(YEAR FROM month)=%s GROUP BY name, target_kes", (YEAR,))
    cur.execute(
        "SELECT name, SUM(target_kes) FROM targets_monthly "
        "WHERE scope='region' AND source='budget' AND EXTRACT(YEAR FROM month)=%s "
        "GROUP BY name ORDER BY name", (YEAR,))
    print("Region budget annual totals:")
    for name, tot in cur.fetchall():
        print(f"  {name:18s} {float(tot):>16,.0f}")
    for month_num, stores in STORE_TARGETS_BY_MONTH.items():
        cur.execute(
            "SELECT SUM(target_kes) FROM targets_monthly "
            "WHERE scope='store' AND source='manual' AND month=%s",
            (date(YEAR, month_num, 1),))
        tot = cur.fetchone()[0] or 0
        print(f"{MONTHS[month_num - 1]} manual per-store total: "
              f"{float(tot):,.0f} across {len(stores)} stores")
    print(f"Upserted {n_region} region rows, {n_store} store rows.")
    cur.close()
    conn.close()


if __name__ == "__main__":
    main()
