"""
Import embedded images from Q3 PD Tracker Excel into pd_style_images.
Each image is anchored at a row; that row's column-A value is the style_number.
Images are resized to max 900px and stored as JPEG base64.
Safe to re-run: upserts on style_id.
"""
import base64
import io
import os
import sys

import openpyxl
import psycopg2
import psycopg2.extras
from PIL import Image as PilImage

EXCEL = "attached_assets/Q3_2026_Product_Development_Tracker_of_New_Styles_1784641645356.xlsx"
DB = os.environ["DATABASE_URL"]

def main():
    print("Loading workbook…")
    wb = openpyxl.load_workbook(EXCEL, data_only=True)
    ws = wb.active

    # Build excel_row → style_number from column A (data starts row 3)
    row_style = {}
    for r in range(3, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v:
            row_style[r] = str(v).strip()

    # Extract image bytes before the workbook state changes
    # anchor._from.row is 0-indexed → excel row = + 1
    image_map = []  # list of (style_number, raw_bytes)
    imgs = list(ws._images)
    print(f"Found {len(imgs)} embedded images")

    for i, img in enumerate(imgs):
        try:
            raw = img._data()
        except Exception as e:
            print(f"  [skip] image {i}: could not read bytes — {e}")
            continue
        ar = img.anchor._from.row
        excel_row = ar + 1
        style_no = row_style.get(excel_row)
        if not style_no:
            print(f"  [skip] image {i}: no style_number at excel row {excel_row}")
            continue
        image_map.append((style_no, raw))

    print(f"Mapped {len(image_map)} images to style numbers")

    conn = psycopg2.connect(DB)
    conn.autocommit = False

    # Fetch style_id by style_number
    with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
        cur.execute("SELECT id, style_number FROM pd_styles WHERE style_number IS NOT NULL")
        rows = cur.fetchall()
    style_id_map = {r["style_number"]: r["id"] for r in rows}
    print(f"pd_styles has {len(style_id_map)} rows with style_number")

    ok = skip_no_style = skip_img_err = already = 0

    with conn.cursor() as cur:
        for style_no, raw in image_map:
            style_id = style_id_map.get(style_no)
            if not style_id:
                print(f"  [skip] {style_no}: not in pd_styles")
                skip_no_style += 1
                continue

            # Resize + convert to JPEG
            try:
                img_obj = PilImage.open(io.BytesIO(raw))
                img_obj = img_obj.convert("RGB")
                img_obj.thumbnail((900, 900), PilImage.LANCZOS)
                buf = io.BytesIO()
                img_obj.save(buf, format="JPEG", quality=82, optimize=True)
                encoded = base64.b64encode(buf.getvalue()).decode()
            except Exception as e:
                print(f"  [skip] {style_no}: image processing error — {e}")
                skip_img_err += 1
                continue

            cur.execute("""
                INSERT INTO pd_style_images
                    (style_id, image_data, content_type, uploaded_by, uploaded_at)
                VALUES (%s, %s, 'image/jpeg', 'excel_import', now())
                ON CONFLICT (style_id) DO UPDATE
                  SET image_data   = EXCLUDED.image_data,
                      content_type = 'image/jpeg',
                      uploaded_by  = 'excel_import',
                      uploaded_at  = now()
            """, (style_id, encoded))
            ok += 1
            print(f"  [ok] {style_no} (id={style_id})")

    conn.commit()
    conn.close()

    print(f"\nDone: {ok} imported, {skip_no_style} style not found, {skip_img_err} image errors")

if __name__ == "__main__":
    main()
